"""PDF and DOCX export for lesson plans and worksheets.

Uses ReportLab (PDF) and python-docx (DOCX). Both take plain data structures,
so they work identically in mock mode and with a real LLM.
"""
from __future__ import annotations

import io
from typing import Any, Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# Devanagari is not in ReportLab's built-in Type1 fonts. We register a Unicode
# TTF when one is available on the machine; otherwise we fall back gracefully
# (the document still builds, and Hindi renders as boxes in the PDF only).
_DEVANAGARI_FONTS = [
    r"C:\Windows\Fonts\Nirmala.ttf",
    r"C:\Windows\Fonts\NirmalaB.ttf",
    r"C:\Windows\Fonts\mangal.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
    "/System/Library/Fonts/Supplemental/DevanagariMT.ttc",
]
_LATIN_FONTS = [
    r"C:\Windows\Fonts\segoeui.ttf",
    r"C:\Windows\Fonts\calibri.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]


def _register_fonts() -> tuple[str, str]:
    """Return (latin_font, hindi_font) names registered with ReportLab."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    latin = "Helvetica"
    for path in _LATIN_FONTS:
        if _exists(path):
            try:
                pdfmetrics.registerFont(TTFont("TC-Latin", path))
                latin = "TC-Latin"
                break
            except Exception:  # noqa: BLE001
                continue

    hindi = latin
    for path in _DEVANAGARI_FONTS:
        if _exists(path):
            try:
                pdfmetrics.registerFont(TTFont("TC-Hindi", path))
                hindi = "TC-Hindi"
                break
            except Exception:  # noqa: BLE001
                continue
    return latin, hindi


def _exists(path: str) -> bool:
    import os

    return os.path.isfile(path)


def _styles(fonts: tuple[str, str]) -> dict[str, ParagraphStyle]:
    latin, hindi = fonts
    base = getSampleStyleSheet()
    title = ParagraphStyle(
        "tcTitle",
        parent=base["Title"],
        fontName=hindi,
        fontSize=18,
        leading=24,
        spaceAfter=4,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1e3a5f"),
    )
    subtitle = ParagraphStyle(
        "tcSubtitle",
        parent=base["Normal"],
        fontName=hindi,
        fontSize=10,
        leading=14,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=14,
    )
    heading = ParagraphStyle(
        "tcHeading",
        parent=base["Heading2"],
        fontName=hindi,
        fontSize=12.5,
        leading=17,
        spaceBefore=12,
        spaceAfter=5,
        textColor=colors.HexColor("#0f766e"),
    )
    body = ParagraphStyle(
        "tcBody", parent=base["Normal"], fontName=hindi, fontSize=10, leading=15
    )
    small = ParagraphStyle(
        "tcSmall",
        parent=base["Normal"],
        fontName=hindi,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#94a3b8"),
    )
    return {"title": title, "subtitle": subtitle, "heading": heading, "body": body, "small": small}


def _esc(text: Any) -> str:
    from xml.sax.saxutils import escape

    return escape(str(text if text is not None else ""))


def _bullets(items: Iterable[Any], style: ParagraphStyle) -> ListFlowable:
    return ListFlowable(
        [ListItem(Paragraph(_esc(i), style), leftIndent=12) for i in items if str(i).strip()],
        bulletType="bullet",
        start="•",
        leftIndent=14,
    )


def _footer(canvas, doc, font: str) -> None:
    canvas.saveState()
    canvas.setFont(font, 8)
    canvas.setFillColor(colors.HexColor("#94a3b8"))
    canvas.drawCentredString(A4[0] / 2, 12 * mm, f"TeacherCopilot  |  Page {doc.page}")
    canvas.restoreState()


# ---------------------------------------------------------------------------
# Lesson plan
# ---------------------------------------------------------------------------
def lesson_plan_to_pdf(plan: dict[str, Any], *, subtitle: str = "") -> bytes:
    fonts = _register_fonts()
    s = _styles(fonts)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=str(plan.get("title", "Lesson Plan")),
        author="TeacherCopilot",
    )
    story: list[Any] = [
        Paragraph(_esc(plan.get("title", "Lesson Plan")), s["title"]),
    ]
    meta_bits = [b for b in [subtitle, "AI-generated draft - review before use"] if b]
    if meta_bits:
        story.append(Paragraph(" &nbsp;|&nbsp; ".join(_esc(b) for b in meta_bits), s["subtitle"]))

    if plan.get("learning_objectives"):
        story += [
            Paragraph("Learning objectives", s["heading"]),
            _bullets(plan["learning_objectives"], s["body"]),
        ]
    if plan.get("prior_knowledge"):
        story += [
            Paragraph("Prior knowledge assumed", s["heading"]),
            _bullets(plan["prior_knowledge"], s["body"]),
        ]
    if plan.get("materials"):
        story += [Paragraph("Materials", s["heading"]), _bullets(plan["materials"], s["body"])]

    if plan.get("lesson_flow"):
        story.append(Paragraph("Lesson flow", s["heading"]))
        rows = [["Phase", "Min", "Teacher actions", "Student actions"]]
        for item in plan["lesson_flow"]:
            rows.append(
                [
                    Paragraph(_esc(item.get("phase", "")).title(), s["body"]),
                    Paragraph(_esc(item.get("minutes", "")), s["body"]),
                    Paragraph(
                        "<br/>".join(_esc(a) for a in item.get("teacher_actions", [])),
                        s["body"],
                    ),
                    Paragraph(
                        "<br/>".join(_esc(a) for a in item.get("student_actions", [])),
                        s["body"],
                    ),
                ]
            )
        total = sum(int(i.get("minutes", 0) or 0) for i in plan["lesson_flow"])
        rows.append(["Total", str(total), "", ""])
        table = Table(rows, colWidths=[22 * mm, 12 * mm, 68 * mm, 68 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTNAME", (0, 0), (-1, 0), fonts[1]),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e2e8f0")),
                ]
            )
        )
        story.append(table)

    if plan.get("key_vocabulary"):
        story += [
            Paragraph("Key vocabulary", s["heading"]),
            _bullets(plan["key_vocabulary"], s["body"]),
        ]
    if plan.get("common_misconceptions"):
        story += [
            Paragraph("Common misconceptions to pre-empt", s["heading"]),
            _bullets(plan["common_misconceptions"], s["body"]),
        ]
    if plan.get("homework"):
        story += [
            Paragraph("Homework", s["heading"]),
            Paragraph(_esc(plan["homework"]), s["body"]),
        ]
    if plan.get("exit_ticket"):
        story += [
            Paragraph("Exit ticket (3 questions)", s["heading"]),
            _bullets(plan["exit_ticket"], s["body"]),
        ]
    if plan.get("assessment_rubric"):
        story.append(Paragraph("Assessment rubric", s["heading"]))
        rows = [["Criterion", "Description", "Marks"]]
        for row in plan["assessment_rubric"]:
            rows.append(
                [
                    Paragraph(_esc(row.get("criterion", "")), s["body"]),
                    Paragraph(_esc(row.get("description", "")), s["body"]),
                    Paragraph(_esc(row.get("marks", "")), s["body"]),
                ]
            )
        table = Table(rows, colWidths=[45 * mm, 105 * mm, 20 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ]
            )
        )
        story.append(table)

    doc.build(story, onFirstPage=lambda c, d: _footer(c, d, fonts[1]),
              onLaterPages=lambda c, d: _footer(c, d, fonts[1]))
    return buffer.getvalue()


def lesson_plan_to_docx(plan: dict[str, Any], *, subtitle: str = "") -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_heading(str(plan.get("title", "Lesson Plan")), level=0)
    if subtitle:
        paragraph = document.add_paragraph(subtitle)
        paragraph.runs[0].font.size = Pt(9)
        paragraph.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    note = document.add_paragraph("AI-generated draft - review before use.")
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(8)
    note.runs[0].font.color.rgb = RGBColor(0x94, 0x3B, 0x8B)

    def add_bullets(items: Iterable[Any]) -> None:
        for item in items:
            if str(item).strip():
                document.add_paragraph(str(item), style="List Bullet")

    if plan.get("learning_objectives"):
        document.add_heading("Learning objectives", level=1)
        add_bullets(plan["learning_objectives"])
    if plan.get("prior_knowledge"):
        document.add_heading("Prior knowledge assumed", level=1)
        add_bullets(plan["prior_knowledge"])
    if plan.get("materials"):
        document.add_heading("Materials", level=1)
        add_bullets(plan["materials"])
    if plan.get("lesson_flow"):
        document.add_heading("Lesson flow", level=1)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["Phase", "Min", "Teacher actions", "Student actions"]):
            cell.text = str(text)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        total = 0
        for item in plan["lesson_flow"]:
            row = table.add_row().cells
            row[0].text = str(item.get("phase", "")).title()
            row[1].text = str(item.get("minutes", ""))
            row[2].text = "; ".join(item.get("teacher_actions", []))
            row[3].text = "; ".join(item.get("student_actions", []))
            total += int(item.get("minutes", 0) or 0)
        row = table.add_row().cells
        row[0].text = "Total"
        row[1].text = str(total)
        for run in row[0].paragraphs[0].runs:
            run.bold = True
    if plan.get("key_vocabulary"):
        document.add_heading("Key vocabulary", level=1)
        add_bullets(plan["key_vocabulary"])
    if plan.get("common_misconceptions"):
        document.add_heading("Common misconceptions to pre-empt", level=1)
        add_bullets(plan["common_misconceptions"])
    if plan.get("homework"):
        document.add_heading("Homework", level=1)
        document.add_paragraph(str(plan["homework"]))
    if plan.get("exit_ticket"):
        document.add_heading("Exit ticket (3 questions)", level=1)
        add_bullets(plan["exit_ticket"])
    if plan.get("assessment_rubric"):
        document.add_heading("Assessment rubric", level=1)
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["Criterion", "Description", "Marks"]):
            cell.text = str(text)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row_data in plan["assessment_rubric"]:
            row = table.add_row().cells
            row[0].text = str(row_data.get("criterion", ""))
            row[1].text = str(row_data.get("description", ""))
            row[2].text = str(row_data.get("marks", ""))

    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Worksheet
# ---------------------------------------------------------------------------
def worksheet_to_pdf(
    level: str,
    material: dict[str, Any],
    *,
    subtitle: str = "",
    include_answers: bool = False,
) -> bytes:
    fonts = _register_fonts()
    s = _styles(fonts)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=f"{level.title()} worksheet",
    )
    label = {"support": "Support", "core": "Core", "extension": "Extension"}.get(level, level)
    story: list[Any] = [
        Paragraph(_esc(f"{label} level worksheet"), s["title"]),
    ]
    meta = [b for b in [subtitle, f"Topic: {material.get('topic', '')}" if material.get("topic") else ""] if b]
    if meta:
        story.append(Paragraph(" &nbsp;|&nbsp; ".join(_esc(b) for b in meta), s["subtitle"]))

    if material.get("what_changed"):
        story += [
            Paragraph("What is different at this level", s["heading"]),
            Paragraph(_esc(material["what_changed"]), s["body"]),
        ]
    if material.get("content"):
        story += [Paragraph("What to learn", s["heading"]), Paragraph(_esc(material["content"]), s["body"])]
    if material.get("key_points"):
        story += [Paragraph("Key points", s["heading"]), _bullets(material["key_points"], s["body"])]
    if material.get("scaffolds"):
        story += [Paragraph("Scaffolds and hints", s["heading"]), _bullets(material["scaffolds"], s["body"])]

    questions = material.get("worksheet", [])
    if questions:
        story.append(Paragraph("Questions", s["heading"]))
        rows = [["#", "Question", "Marks", "Hint"]]
        for index, question in enumerate(questions, 1):
            hint = question.get("hint", "")
            rows.append(
                [
                    Paragraph(str(index), s["body"]),
                    Paragraph(_esc(question.get("question", "")), s["body"]),
                    Paragraph(_esc(question.get("marks", "")), s["body"]),
                    Paragraph(_esc(hint), s["small"]),
                ]
            )
        table = Table(rows, colWidths=[10 * mm, 100 * mm, 15 * mm, 45 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ]
            )
        )
        story.append(table)

    if include_answers and material.get("answer_key"):
        story.append(PageBreak())
        story.append(Paragraph("Answer key (teacher copy)", s["heading"]))
        rows = [["#", "Question", "Model answer", "Marks"]]
        key_by_id = {str(a.get("question_id")): a for a in material["answer_key"]}
        for index, question in enumerate(questions, 1):
            answer = key_by_id.get(str(question.get("id")), {})
            rows.append(
                [
                    Paragraph(str(index), s["body"]),
                    Paragraph(_esc(question.get("question", "")), s["small"]),
                    Paragraph(_esc(answer.get("answer", "")), s["body"]),
                    Paragraph(_esc(answer.get("marks", "")), s["body"]),
                ]
            )
        table = Table(rows, colWidths=[10 * mm, 55 * mm, 90 * mm, 15 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#475569")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(table)

    doc.build(story, onFirstPage=lambda c, d: _footer(c, d, fonts[1]),
              onLaterPages=lambda c, d: _footer(c, d, fonts[1]))
    return buffer.getvalue()


def worksheet_to_docx(
    level: str,
    material: dict[str, Any],
    *,
    subtitle: str = "",
    include_answers: bool = False,
) -> bytes:
    buffer = io.BytesIO()
    document = Document()
    label = {"support": "Support", "core": "Core", "extension": "Extension"}.get(level, level)
    document.add_heading(f"{label} level worksheet", level=0)
    if subtitle:
        paragraph = document.add_paragraph(subtitle)
        paragraph.runs[0].font.size = Pt(9)
    note = document.add_paragraph("AI-generated draft - review before use.")
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(8)

    if material.get("what_changed"):
        document.add_heading("What is different at this level", level=1)
        document.add_paragraph(str(material["what_changed"]))
    if material.get("content"):
        document.add_heading("What to learn", level=1)
        document.add_paragraph(str(material["content"]))
    for heading, key in (("Key points", "key_points"), ("Scaffolds and hints", "scaffolds")):
        if material.get(key):
            document.add_heading(heading, level=1)
            for item in material[key]:
                if str(item).strip():
                    document.add_paragraph(str(item), style="List Bullet")

    questions = material.get("worksheet", [])
    if questions:
        document.add_heading("Questions", level=1)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["#", "Question", "Marks", "Hint"]):
            cell.text = text
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for index, question in enumerate(questions, 1):
            row = table.add_row().cells
            row[0].text = str(index)
            row[1].text = str(question.get("question", ""))
            row[2].text = str(question.get("marks", ""))
            row[3].text = str(question.get("hint", ""))

    if include_answers and material.get("answer_key"):
        document.add_page_break()
        document.add_heading("Answer key (teacher copy)", level=1)
        key_by_id = {str(a.get("question_id")): a for a in material["answer_key"]}
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["#", "Question", "Model answer"]):
            cell.text = text
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for index, question in enumerate(questions, 1):
            answer = key_by_id.get(str(question.get("id")), {})
            row = table.add_row().cells
            row[0].text = str(index)
            row[1].text = str(question.get("question", ""))
            row[2].text = str(answer.get("answer", ""))

    document.save(buffer)
    return buffer.getvalue()
